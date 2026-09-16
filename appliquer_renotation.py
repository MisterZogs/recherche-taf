"""
Applique les résultats d'un lot de renotation (agents du 16/09/2026) au
tableur : met à jour Priorité / Fit-Notes (avec marqueur [revu ...]) pour
les offres vivantes, marque Statut=Expiré pour les liens morts (ce qui les
archivera automatiquement vers Fait au prochain passage d'ajouter_offres),
puis retrie via ajouter_offres([]).

Usage : python3 appliquer_renotation.py <fichier_resultats.json> [<fichier2.json> ...]
"""
import json
import sys
from datetime import date

import openpyxl

from add_offre import ajouter_offres, FICHIER

MARQUEUR = f"[revu {date.today().isoformat()}]"


def appliquer(fichiers):
    wb = openpyxl.load_workbook(FICHIER)
    tabs = ['Offres SIRH', 'Pays Basque', 'Offres USA', 'Offres PM',
            'Offres UX', 'Offres SEO', 'Offres CSM', 'Offres IA']

    # Index Lien -> (feuille, ligne) sur les onglets actifs
    lien_index = {}
    col_cache = {}
    for sn in tabs:
        ws = wb[sn]
        headers = [c.value for c in ws[1]]
        idx = {h: i + 1 for i, h in enumerate(headers)}
        col_cache[sn] = idx
        for r in range(2, ws.max_row + 1):
            v = ws.cell(row=r, column=idx['Lien']).value
            if v:
                lien_index[str(v).strip()] = (sn, r)

    maj = expire = introuvable = deja_revu = 0
    for fichier in fichiers:
        with open(fichier) as f:
            resultats = json.load(f)
        for res in resultats:
            lien = str(res.get('Lien') or '').strip()
            if not lien or lien not in lien_index:
                introuvable += 1
                print(f"  [introuvable] {lien}")
                continue
            sn, r = lien_index[lien]
            idx = col_cache[sn]
            fit_cell = ws = wb[sn].cell(row=r, column=idx['Fit / Notes'])
            if fit_cell.value and MARQUEUR in str(fit_cell.value):
                deja_revu += 1
                continue
            if res.get('statut') == 'Expiré':
                wb[sn].cell(row=r, column=idx['Statut']).value = 'Expiré'
                note = res.get('nouveau_fit') or 'Lien mort constaté lors de la repasse de renotation'
                fit_cell.value = f"{MARQUEUR} {note}"
                expire += 1
            else:
                nv_prio = res.get('nouvelle_priorite')
                nv_fit = res.get('nouveau_fit')
                if nv_prio:
                    wb[sn].cell(row=r, column=idx['Priorité']).value = nv_prio
                if nv_fit:
                    fit_cell.value = f"{MARQUEUR} {nv_fit}"
                else:
                    fit_cell.value = f"{MARQUEUR} {fit_cell.value or ''}"
                maj += 1

    wb.save(FICHIER)
    print(f"\nMises à jour : {maj} | Expirées : {expire} | Déjà revues (ignorées) : {deja_revu} | Introuvables : {introuvable}")

    # Retri + archivage des lignes Expiré vers Fait
    ajouter_offres([], verbose=False)
    print("Retri + archivage Fait effectué.")


if __name__ == '__main__':
    appliquer(sys.argv[1:])
